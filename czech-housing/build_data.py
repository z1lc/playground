# /// script
# requires-python = ">=3.11"
# dependencies = ["httpx", "openpyxl"]
# ///
"""Build the Czech housing dataset.

    uv run build_data.py            # fetch (cached), validate, write data.json, inline into index.html
    uv run build_data.py --refresh  # ignore the cache and re-download everything
    uv run build_data.py --offline  # fail on a cache miss instead of fetching

Every series comes straight from its publisher. The one hand-entered figure is the 2011
census dwelling count that anchors the Czech housing-stock series.

Sources, and what each one feeds:
  * EMF Hypostat 2026      dwelling stock and completions for the peer countries.
  * CZSO DataStat          Czech completions (STA09AT1), used to build the Czech stock
                           series from the 2011 census, because Hypostat has no Czech
                           stock figures for 2014-2021 and the 2021 census cannot be
                           spliced on (it counted ~290k more dwellings than the 2011
                           census plus a decade of completions).
  * Eurostat demo_gind     1 January population and natural change.
  * Eurostat ilc_lvps08    share of 25-34-year-olds living with their parents.
  * OECD house prices      rent index (RPI), nominal (HPI) and real (RHP) house prices.
                           Real rents are RPI deflated by the same consumption deflator
                           the OECD uses for RHP, i.e. RPI * RHP / HPI.
  * World Bank DB2020      days to obtain a construction permit (Doing Business 2020,
                           data year 2019; the series was discontinued in 2021).

Two cleaning steps are applied and reported in the page footnotes:
  * Hypostat stock series have breaks (definition changes, one-off typos). A one-year
    outlier is dropped and interpolated; a persistent jump outside -0.5%..+2.5% a year
    is bridged with the country's median growth. More than two problems drops the country.
  * Census revisions show up as population jumps that natural change and migration do
    not explain. Each is spread back linearly over the preceding decade, which is how
    statistical offices rebuild intercensal estimates.
"""

import argparse
import csv
import io
import json
import re
import statistics
import sys
import time
from datetime import date
from pathlib import Path
from typing import Any

import httpx
import openpyxl

HERE = Path(__file__).resolve().parent
RAW = HERE / "raw_responses"
DATA_JSON = HERE / "data.json"
INDEX_HTML = HERE / "index.html"

FOCUS = "CZ"
LABELLED = ["PL", "SK", "AT", "DE"]

# id, Eurostat geo, ISO3 (OECD / World Bank), Hypostat row name, display name
COUNTRY_ROWS = [
    ("AT", "AT", "AUT", "Austria", "Austria"),
    ("BE", "BE", "BEL", "Belgium", "Belgium"),
    ("CH", "CH", "CHE", "Switzerland", "Switzerland"),
    ("CZ", "CZ", "CZE", "Czechia", "Czechia"),
    ("DE", "DE", "DEU", "Germany", "Germany"),
    ("DK", "DK", "DNK", "Denmark", "Denmark"),
    ("ES", "ES", "ESP", "Spain", "Spain"),
    ("FI", "FI", "FIN", "Finland", "Finland"),
    ("FR", "FR", "FRA", "France", "France"),
    ("GB", "UK", "GBR", "United Kingdom", "UK"),
    ("IE", "IE", "IRL", "Ireland", "Ireland"),
    ("IT", "IT", "ITA", "Italy", "Italy"),
    ("NL", "NL", "NLD", "Netherlands", "Netherlands"),
    ("NO", "NO", "NOR", "Norway", "Norway"),
    ("PL", "PL", "POL", "Poland", "Poland"),
    ("PT", "PT", "PRT", "Portugal", "Portugal"),
    ("SE", "SE", "SWE", "Sweden", "Sweden"),
    ("SK", "SK", "SVK", "Slovakia", "Slovakia"),
]
COUNTRIES = [dict(zip(("id", "eurostat", "iso3", "hypostat", "name"), row)) for row in COUNTRY_ROWS]
BY_ID = {c["id"]: c for c in COUNTRIES}

PARENTS_COUNTRIES = ["CZ", "SK", "PL", "AT", "DE", "ES"]

BASE_YEAR = 2015
PANEL_YEARS = list(range(2010, 2026))
PARENTS_YEARS = list(range(2005, 2026))
CZ_YEARS = list(range(2005, 2027))
COMPLETIONS_YEARS = list(range(2015, 2025))

# 2011 census (SLDB, 26 March 2011): all dwellings, occupied and unoccupied.
CENSUS_2011_DWELLINGS = 4_756_572
# CZSO, "Domovní a bytový fond podle výsledků SLDB" (2014), hosted by the National Repository of Grey Literature.
CENSUS_2011_URL = "https://invenio.nusl.cz/record/204140/files/nusl-204140_1.pdf"

EUROSTAT = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data"
OECD_HOUSE = (
    "https://sdmx.oecd.org/public/rest/data/OECD.ECO.MPD,DSD_AN_HOUSE_PRICES@DF_HOUSE_PRICES,1.0/"
    "{areas}.A.RPI+RHP+HPI.?format=csv&startPeriod=2005"
)
HYPOSTAT = (
    "https://hypo.org/sites/default/files/2026-09/Final-Statistical-Tables-Hypostat-2026-to-share_UPDATE-14-09.xlsx"
)
CZSO_COMPLETIONS = "https://data.csu.gov.cz/api/dotaz/v1/data/vybery/STA09AT1?format=CSV"
DB2020 = (
    "https://archive.doingbusiness.org/content/dam/doingBusiness/excel/db2020/"
    "Historical-data---COMPLETE-dataset-with-scores.xlsx"
)

SOURCES = {
    "hypostat": {
        "name": "European Mortgage Federation, Hypostat 2026",
        "url": "https://hypo.org/ecbc/publications/hypostat/",
    },
    "czso": {
        "name": "Czech Statistical Office, completed dwellings (DataStat STA09AT1)",
        "url": "https://csu.gov.cz/produkty/dokoncene-byty-v-obcich",
    },
    "census": {"name": "Czech Statistical Office, housing stock in the 2011 census", "url": CENSUS_2011_URL},
    "demo": {
        "name": "Eurostat, population change (demo_gind)",
        "url": "https://ec.europa.eu/eurostat/databrowser/view/demo_gind/default/table",
    },
    "parents": {
        "name": "Eurostat, young people living with their parents (ilc_lvps08)",
        "url": "https://ec.europa.eu/eurostat/databrowser/view/ilc_lvps08/default/table",
    },
    "oecd": {
        "name": "OECD, Analytical house price indicators",
        "url": "https://data-explorer.oecd.org/vis?df[ds]=dsDisseminateFinalDMZ&df[id]=DSD_AN_HOUSE_PRICES%40DF_HOUSE_PRICES&df[ag]=OECD.ECO.MPD",
    },
    "db": {
        "name": "World Bank, Doing Business 2020 (dealing with construction permits)",
        "url": "https://archive.doingbusiness.org/en/data/exploretopics/dealing-with-construction-permits",
    },
}


class Fetcher:
    def __init__(self, refresh: bool, offline: bool) -> None:
        self.refresh = refresh
        self.offline = offline
        self.client = httpx.Client(
            timeout=300, follow_redirects=True, headers={"User-Agent": "playground-czech-housing/1.0"}
        )

    def get(self, name: str, url: str) -> bytes:
        path = RAW / name
        if path.exists() and not self.refresh:
            return path.read_bytes()
        if self.offline:
            sys.exit(f"cache miss in --offline mode: {name}")
        for attempt in range(4):
            print(f"  fetching {name} ...", file=sys.stderr)
            r = self.client.get(url)
            if r.status_code in (429, 502, 503) and attempt < 3:
                time.sleep(10 * (attempt + 1))
                continue
            r.raise_for_status()
            RAW.mkdir(exist_ok=True)
            path.write_bytes(r.content)
            return r.content
        raise RuntimeError(f"giving up on {url}")


# ---------------------------------------------------------------- parsing helpers


def jsonstat_rows(j: dict[str, Any]) -> list[tuple[dict[str, str], float]]:
    """Flatten a Eurostat JSON-stat 2.0 dataset into (coordinates, value) pairs."""
    dims, sizes = j["id"], j["size"]
    cats = [sorted(j["dimension"][d]["category"]["index"].items(), key=lambda kv: kv[1]) for d in dims]
    out = []
    for flat, value in j["value"].items():
        n, coords = int(flat), {}
        for d, s, c in reversed(list(zip(dims, sizes, cats))):
            coords[d] = c[n % s][0]
            n //= s
        out.append((coords, float(value)))
    return out


def as_year(cell: Any) -> int | None:
    try:
        y = int(float(str(cell).strip()))
    except (TypeError, ValueError):
        return None
    return y if 1950 <= y <= 2100 else None


def hypostat_sheet(wb: Any, sheet: str) -> dict[str, dict[int, float]]:
    rows = list(wb[sheet].iter_rows(values_only=True))
    header = next(r for r in rows if sum(as_year(c) is not None for c in r[1:]) > 5)
    years = {j: as_year(c) for j, c in enumerate(header) if j > 0 and as_year(c) is not None}
    out: dict[str, dict[int, float]] = {}
    for r in rows:
        if not isinstance(r[0], str):
            continue
        name = re.sub(r"[*\d]", "", r[0]).strip()
        vals = {years[j]: float(r[j]) for j in years if j < len(r) and isinstance(r[j], (int, float))}
        if vals:
            out[name] = vals
    return out


def rnd(v: float | None, nd: int = 2) -> float | None:
    return None if v is None else round(v, nd)


def index_to_base(series: dict[int, float], years: list[int]) -> list[float | None] | None:
    if BASE_YEAR not in series:
        return None
    base = series[BASE_YEAR]
    return [rnd(series[y] / base * 100) if y in series else None for y in years]


def join_names(ids: list[str]) -> str:
    names = [BY_ID[i]["name"] for i in ids]
    return names[0] if len(names) == 1 else ", ".join(names[:-1]) + " and " + names[-1]


# ---------------------------------------------------------------- loaders


def load_hypostat(fetcher: Fetcher) -> tuple[dict[str, dict[int, float]], dict[str, dict[int, float]]]:
    wb = openpyxl.load_workbook(io.BytesIO(fetcher.get("hypostat_2026.xlsx", HYPOSTAT)), read_only=True, data_only=True)
    stock_rows, compl_rows = hypostat_sheet(wb, "16. TDwe Stk"), hypostat_sheet(wb, "14. House Com")
    stock = {c["id"]: stock_rows.get(c["hypostat"], {}) for c in COUNTRIES}
    completions = {c["id"]: compl_rows.get(c["hypostat"], {}) for c in COUNTRIES}
    return stock, completions


def load_population(fetcher: Fetcher) -> dict[str, dict[str, dict[int, float]]]:
    geos = "".join(f"&geo={c['eurostat']}" for c in COUNTRIES)
    url = f"{EUROSTAT}/demo_gind?indic_de=JAN&indic_de=NATGROW&indic_de=CNMIGRAT&sinceTimePeriod=2000{geos}"
    j = json.loads(fetcher.get("eurostat_demo_gind.json", url))
    geo_to_id = {c["eurostat"]: c["id"] for c in COUNTRIES}
    out: dict[str, dict[str, dict[int, float]]] = {
        c["id"]: {"JAN": {}, "NATGROW": {}, "CNMIGRAT": {}} for c in COUNTRIES
    }
    for coords, v in jsonstat_rows(j):
        out[geo_to_id[coords["geo"]]][coords["indic_de"]][int(coords["time"])] = v
    return out


def load_parents(fetcher: Fetcher) -> dict[str, dict[int, float]]:
    geos = "".join(f"&geo={g}" for g in PARENTS_COUNTRIES + ["EU27_2020"])
    url = f"{EUROSTAT}/ilc_lvps08?age=Y25-34&sex=T&unit=PC{geos}"
    j = json.loads(fetcher.get("eurostat_ilc_lvps08.json", url))
    out: dict[str, dict[int, float]] = {}
    for coords, v in jsonstat_rows(j):
        out.setdefault(coords["geo"], {})[int(coords["time"])] = v
    return out


def load_oecd(fetcher: Fetcher) -> dict[str, dict[str, dict[int, float]]]:
    url = OECD_HOUSE.format(areas="+".join(c["iso3"] for c in COUNTRIES))
    text = fetcher.get("oecd_house_prices.csv", url).decode("utf-8-sig")
    iso_to_id = {c["iso3"]: c["id"] for c in COUNTRIES}
    out: dict[str, dict[str, dict[int, float]]] = {c["id"]: {"RPI": {}, "RHP": {}, "HPI": {}} for c in COUNTRIES}
    for r in csv.DictReader(io.StringIO(text)):
        if r["REF_AREA"] in iso_to_id and r["OBS_VALUE"]:
            out[iso_to_id[r["REF_AREA"]]][r["MEASURE"]][int(r["TIME_PERIOD"])] = float(r["OBS_VALUE"])
    return out


def load_czso_completions(fetcher: Fetcher) -> dict[int, float]:
    text = fetcher.get("czso_sta09at1.csv", CZSO_COMPLETIONS).decode("utf-8-sig")
    return {
        int(r["Roky"]): float(r["Hodnota"])
        for r in csv.DictReader(io.StringIO(text))
        if r["Ukazatel"] == "Dokončené byty" and r["Hodnota"]
    }


def load_permit_days(fetcher: Fetcher) -> dict[str, float]:
    wb = openpyxl.load_workbook(
        io.BytesIO(fetcher.get("db2020_historical.xlsx", DB2020)), read_only=True, data_only=True
    )
    rows = list(wb["All Data"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] == "Country code")
    header = rows[hi]
    start = header.index("Rank-Dealing with construction permits")
    col = next(j for j in range(start, len(header)) if header[j] == "Time (days)")
    iso_to_id = {c["iso3"]: c["id"] for c in COUNTRIES}
    out = {}
    for r in rows[hi + 1 :]:
        if r[0] in iso_to_id and as_year(r[4]) == 2020:
            try:
                out[iso_to_id[r[0]]] = float(str(r[col]).strip())  # the workbook stores numbers as text
            except ValueError:
                pass
    return out


# ---------------------------------------------------------------- cleaning


STOCK_LO, STOCK_HI = -0.005, 0.025


def plausible(a: float, b: float, years: int = 1) -> bool:
    return STOCK_LO * years <= b / a - 1 <= STOCK_HI * years


def clean_stock(raw: dict[int, float], years: list[int]) -> tuple[dict[int, float] | None, list[str]]:
    """Repair a Hypostat dwelling-stock series over `years`. Returns (series, fixes) or (None, reason)."""
    v: dict[int, float | None] = {y: raw.get(y) for y in years}
    fixes: list[str] = []
    for a, b, c in zip(years, years[1:], years[2:]):
        va, vb, vc = v[a], v[b], v[c]
        if None not in (va, vb, vc) and not plausible(va, vb) and not plausible(vb, vc) and plausible(va, vc, 2):
            v[b] = None
            fixes.append(f"{b} outlier")
    known = [y for y in years if v[y] is not None]
    if not known:
        return None, ["no data"]
    for a, b in zip(known, known[1:]):
        gap = b - a - 1
        if gap > 2:
            return None, [f"no data {a + 1}–{b - 1}"]
        if gap:
            r = (v[b] / v[a]) ** (1 / (b - a))
            for k in range(1, gap + 1):
                v[a + k] = v[a] * r**k
            fixes.append(f"{a + 1} gap" if gap == 1 else f"{a + 1}–{b - 1} gap")
    span = list(range(known[0], known[-1] + 1))
    growth = {b: v[b] / v[a] - 1 for a, b in zip(span, span[1:])}
    bad = [y for y, g in growth.items() if not STOCK_LO <= g <= STOCK_HI]
    if len(fixes) + len(bad) > 2:
        return None, [f"{len(fixes) + len(bad)} series breaks"]
    if bad:
        med = statistics.median(g for y, g in growth.items() if y not in bad)
        for y in bad:
            growth[y] = med
            fixes.append(f"{y} break")
    chained = {span[0]: v[span[0]]}
    for y in span[1:]:
        chained[y] = chained[y - 1] * (1 + growth[y])
    return chained, fixes


def smooth_census_breaks(pop: dict[str, dict[int, float]]) -> tuple[dict[int, float], list[tuple[int, float]]]:
    """Spread unexplained population jumps (census revisions) back over the preceding decade."""
    jan, nat, mig = pop["JAN"], pop["NATGROW"], pop["CNMIGRAT"]
    out = dict(jan)
    breaks = []
    for y in sorted(jan):
        if all(y - 1 in s for s in (jan, nat, mig)):
            resid = jan[y] - jan[y - 1] - nat[y - 1] - mig[y - 1]
            if abs(resid) > 0.002 * jan[y]:
                breaks.append((y, resid))
                for k in range(1, 10):
                    if y - 10 + k in out:
                        out[y - 10 + k] += resid * k / 10
    return out, breaks


def build_cz_stock(completions: dict[int, float]) -> dict[int, float]:
    """End-of-year Czech dwelling stock: the March 2011 census, plus or minus completions since."""
    stock = {2010: float(CENSUS_2011_DWELLINGS)}
    for y in range(2011, max(completions) + 1):
        stock[y] = stock[y - 1] + completions[y]
    for y in range(2010, 2003, -1):
        stock[y - 1] = stock[y] - completions[y]
    return stock


# ---------------------------------------------------------------- assembly


def panel(pid: str, title: str, unit: str, series: dict[str, dict[int, float]], notes: list[str]) -> dict[str, Any]:
    out, missing = {}, []
    for c in COUNTRIES:
        s = series.get(c["id"])
        idx = index_to_base(s, PANEL_YEARS) if s else None
        if idx is None:
            missing.append(c["id"])
        else:
            out[c["id"]] = idx
    if missing:
        notes = [f"{join_names(missing)} {'is' if len(missing) == 1 else 'are'} omitted for lack of data."] + notes
    return {"id": pid, "title": title, "unit": unit, "series": out, "notes": notes}


def build(fetcher: Fetcher) -> tuple[dict[str, Any], dict[str, Any]]:
    stock_raw, compl_raw = load_hypostat(fetcher)
    pop = load_population(fetcher)
    parents = load_parents(fetcher)
    oecd = load_oecd(fetcher)
    cz_compl = load_czso_completions(fetcher)
    permits = load_permit_days(fetcher)

    # Chart 1, panel 1: housing stock.
    cz_stock = build_cz_stock(cz_compl)
    stock, stock_fixes = {}, []
    for c in COUNTRIES:
        if c["id"] == FOCUS:
            stock[FOCUS] = cz_stock
            continue
        cleaned, fixes = clean_stock(stock_raw[c["id"]], PANEL_YEARS)
        if cleaned is not None:
            stock[c["id"]] = cleaned
            if fixes:
                stock_fixes.append(f"{c['name']} ({', '.join(fixes)})")
    stock_notes = ["Czechia is the 2011 census count plus dwellings completed since, without subtracting demolitions."]
    if stock_fixes:
        stock_notes.append("Series breaks are bridged with typical growth: " + "; ".join(stock_fixes) + ".")

    # Chart 1, panel 2: population, with census revisions smoothed.
    pop_adj, census_breaks = {}, {}
    for c in COUNTRIES:
        pop_adj[c["id"]], census_breaks[c["id"]] = smooth_census_breaks(pop[c["id"]])
    in_window = [
        f"{BY_ID[i]['name']} {y}"
        for i, bs in census_breaks.items()
        for y, _ in bs
        if PANEL_YEARS[0] < y <= PANEL_YEARS[-1]
    ]
    pop_notes = ["Population on 1 January."]
    if in_window:
        pop_notes.append("Census revisions (" + ", ".join(in_window) + ") are spread back over the preceding decade.")

    # Chart 1, panels 3 and 4: real rents and real house prices.
    real_rents, real_prices = {}, {}
    for c in COUNTRIES:
        o = oecd[c["id"]]
        real_rents[c["id"]] = {
            y: o["RPI"][y] * o["RHP"][y] / o["HPI"][y] for y in o["RHP"] if y in o["RPI"] and y in o["HPI"]
        }
        real_prices[c["id"]] = o["RHP"]
    rent_notes = [
        "Rent indices are deflated by consumer prices. Czechia's official index covers all tenancies, "
        "including municipal flats and leases signed long ago, so it rises more slowly than asking rents. "
        "Controlled rents on older leases were phased out between 2007 and 2012."
    ]

    chart1 = {
        "years": PANEL_YEARS,
        "panels": [
            panel("stock", "Housing stock", "index", stock, stock_notes),
            panel("population", "Population", "index", pop_adj, pop_notes),
            panel("rents", "Real rents", "index", real_rents, rent_notes),
            panel("prices", "Real house prices", "index", real_prices, ["House prices deflated by consumer prices."]),
        ],
    }

    # Chart 2: living with parents.
    chart2 = {
        "years": PARENTS_YEARS,
        "countries": PARENTS_COUNTRIES,
        "series": {g: [rnd(parents.get(g, {}).get(y), 1) for y in PARENTS_YEARS] for g in PARENTS_COUNTRIES},
        "eu": [rnd(parents.get("EU27_2020", {}).get(y), 1) for y in PARENTS_YEARS],
    }

    # Chart 3: Czech stock vs population vs population without migration, all at 1 January.
    cz = pop[FOCUS]
    cz_pop = pop_adj[FOCUS]
    base_pop = cz_pop[BASE_YEAR]
    natural = {BASE_YEAR: base_pop}
    for y in range(BASE_YEAR + 1, CZ_YEARS[-1] + 1):
        natural[y] = natural[y - 1] + cz["NATGROW"][y - 1]
    for y in range(BASE_YEAR - 1, CZ_YEARS[0] - 1, -1):
        natural[y] = natural[y + 1] - cz["NATGROW"][y]
    stock_jan = {y: cz_stock[y - 1] for y in CZ_YEARS if y - 1 in cz_stock}
    cz_break = dict(census_breaks[FOCUS])
    chart3 = {
        "years": CZ_YEARS,
        "population": index_to_base(cz_pop, CZ_YEARS),
        "natural": index_to_base(natural, CZ_YEARS),
        "stock": index_to_base(stock_jan, CZ_YEARS),
        "census2021": round(cz_break.get(2021, 0)),
        "refugeeJump": round(cz["JAN"][2023] - cz["JAN"][2022]),
        "naturalChange2025": round(cz["NATGROW"][2025]),
    }

    # Chart 4: completions per 1,000 residents and permit days.
    completions = []
    for c in COUNTRIES:
        src = cz_compl if c["id"] == FOCUS else compl_raw[c["id"]]
        rates = [
            src[y] / pop[c["id"]]["JAN"][y] * 1000 for y in COMPLETIONS_YEARS if y in src and y in pop[c["id"]]["JAN"]
        ]
        if len(rates) >= 5:
            completions.append({"id": c["id"], "value": rnd(sum(rates) / len(rates)), "years": len(rates)})
    completions.sort(key=lambda d: -d["value"])
    partial = [d["id"] for d in completions if d["years"] < len(COMPLETIONS_YEARS)]
    compl_missing = [c["id"] for c in COUNTRIES if c["id"] not in {d["id"] for d in completions}]
    permit_rows = sorted(({"id": i, "value": v} for i, v in permits.items()), key=lambda d: -d["value"])
    chart4 = {
        "completions": completions,
        "completionsNotes": (
            ([f"{join_names(partial)}: average of the years available."] if partial else [])
            + ([f"{join_names(compl_missing)} omitted for lack of data."] if compl_missing else [])
        ),
        "permits": permit_rows,
    }

    payload = {
        "meta": {"built": date.today().isoformat(), "focus": FOCUS, "labelled": LABELLED, "sources": SOURCES},
        "countries": {c["id"]: c["name"] for c in COUNTRIES},
        "chart1": chart1,
        "chart2": chart2,
        "chart3": chart3,
        "chart4": chart4,
    }
    checks = {
        "cz_pop_2015_raw": cz["JAN"][2015],
        "cz_census_resid_2021": cz_break.get(2021, 0),
        "cz_parents_2015": parents["CZ"][2015],
        "cz_parents_2025": parents["CZ"][2025],
        "cz_completions_2024": cz_compl[2024],
        "cz_completions_per_1000_2024": cz_compl[2024] / cz["JAN"][2024] * 1000,
        "cz_real_rent_2025": real_rents[FOCUS][2025] / real_rents[FOCUS][BASE_YEAR] * 100,
        "cz_real_price_2025": real_prices[FOCUS][2025] / real_prices[FOCUS][BASE_YEAR] * 100,
        "cz_permit_days": permits.get(FOCUS),
    }
    return payload, checks


ANCHORS = {
    "cz_pop_2015_raw": (10_538_275, 10_538_275),
    "cz_census_resid_2021": (-220_000, -195_000),
    "cz_parents_2015": (34.4, 34.4),
    "cz_parents_2025": (22.3, 22.3),
    "cz_completions_2024": (30_274, 30_274),
    "cz_completions_per_1000_2024": (2.6, 3.0),
    "cz_real_rent_2025": (92, 99),
    "cz_real_price_2025": (150, 165),
    "cz_permit_days": (246, 246),
}


def validate(checks: dict[str, Any]) -> list[str]:
    failures = []
    for k, (lo, hi) in ANCHORS.items():
        v = checks.get(k)
        if v is None or not lo - 1e-9 <= v <= hi + 1e-9:
            failures.append(f"{k} = {v}, expected {lo}..{hi}")
    return failures


def inline_into_html(payload: dict[str, Any]) -> bool:
    if not INDEX_HTML.exists():
        print(f"  WARNING: {INDEX_HTML} does not exist yet - skipping inline", file=sys.stderr)
        return False
    text = INDEX_HTML.read_text()
    blob = json.dumps(payload, separators=(",", ":"), ensure_ascii=True).replace("<", "\\u003c")
    js = "    const DATA = " + blob + ";\n"
    pattern = re.compile(r"(// DATA_START\n)(.*?)(    // DATA_END)", re.DOTALL)
    if not pattern.search(text):
        print(f"  WARNING: no DATA_START/DATA_END sentinels in {INDEX_HTML}", file=sys.stderr)
        return False
    INDEX_HTML.write_text(pattern.sub(lambda m: m.group(1) + js + m.group(3), text))
    return True


def summarize(payload: dict[str, Any], checks: dict[str, Any]) -> None:
    c1 = payload["chart1"]
    print("\nChart 1 (2015 = 100), Czechia vs peer median in 2025:")
    for p in c1["panels"]:
        last = c1["years"].index(2025)
        vals = [s[last] for k, s in p["series"].items() if k != FOCUS and s[last] is not None]
        cz = p["series"].get(FOCUS, [None] * len(c1["years"]))[last]
        print(f"  {p['title']:<18} CZ {cz:>7}  peers median {statistics.median(vals):>7.1f}  n={len(p['series'])}")
        for n in p["notes"]:
            print(f"      note: {n}")
    c2 = payload["chart2"]
    print("\nChart 2 (% aged 25-34 living with parents), 2015 -> 2025:")
    for g in c2["countries"]:
        s = dict(zip(c2["years"], c2["series"][g]))
        print(f"  {g}: {s.get(2015)} -> {s.get(2025)}")
    c3 = payload["chart3"]
    print("\nChart 3 (Czechia, 1 Jan, 2015 = 100):")
    for y in (2005, 2010, 2015, 2020, 2021, 2022, 2023, 2026):
        i = c3["years"].index(y)
        print(f"  {y}: population {c3['population'][i]}  no-migration {c3['natural'][i]}  stock {c3['stock'][i]}")
    print(
        f"  census 2021 revision {c3['census2021']:+,}; 2022 jump {c3['refugeeJump']:+,}; natural change 2025 {c3['naturalChange2025']:+,}"
    )
    c4 = payload["chart4"]
    print("\nChart 4: completions per 1,000 (2015-2024 avg) | permit days (DB2020)")
    for d in c4["completions"]:
        print(f"  {d['id']} {d['value']:>5}  ({d['years']} yrs)")
    print("  " + ", ".join(f"{d['id']} {d['value']:.0f}" for d in c4["permits"]))
    for n in c4["completionsNotes"]:
        print(f"      note: {n}")
    print("\nAnchors:")
    for k, v in checks.items():
        print(f"  {k}: {v:,.2f}" if isinstance(v, float) else f"  {k}: {v}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--refresh", action="store_true", help="ignore the cache and re-download")
    ap.add_argument("--offline", action="store_true", help="fail on a cache miss instead of fetching")
    args = ap.parse_args()
    payload, checks = build(Fetcher(args.refresh, args.offline))
    summarize(payload, checks)
    failures = validate(checks)
    if failures:
        print("\nVALIDATION FAILED:\n  " + "\n  ".join(failures), file=sys.stderr)
        sys.exit(1)
    DATA_JSON.write_text(json.dumps(payload, indent=1, ensure_ascii=False) + "\n")
    inlined = inline_into_html(payload)
    print(f"\nwrote {DATA_JSON.name}" + (" and inlined into index.html" if inlined else ""))


if __name__ == "__main__":
    main()
