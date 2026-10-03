# /// script
# requires-python = ">=3.11"
# dependencies = ["httpx", "openpyxl", "xlrd"]
# ///
"""Build the Housing Squeeze dataset: Czechia and the United States against rich-country peers.

    uv run build_data.py            # fetch (cached), validate, write data.json, inline into index.html
    uv run build_data.py --refresh  # ignore the cache and re-download everything
    uv run build_data.py --offline  # fail on a cache miss instead of fetching

Every annual series comes straight from its publisher. Hand-entered figures live in MANUAL
with a source URL each: census and survey counts that exist only as publications.

Sources, and what each one feeds:
  * EMF Hypostat 2026      dwelling stock and completions for the European peers, Canada
                           and Australia (completions only).
  * CZSO DataStat          Czech completions (STA09AT1), used to build the Czech stock
                           series from the 2011 census, because Hypostat has no Czech
                           stock figures for 2014-2021 and the 2021 census cannot be
                           spliced on (it counted ~290k more dwellings than the 2011
                           census plus a decade of completions).
  * Eurostat demo_gind     1 January population and natural change, Europe.
  * Eurostat ilc_lvps08    share of 25-34-year-olds living with their parents, Europe.
  * OECD population        mid-year population for the US, Canada, Australia and Japan.
  * OECD house prices      rent index (RPI), nominal (HPI) and real (RHP) house prices.
                           Real rents are RPI deflated by the same consumption deflator
                           the OECD uses for RHP, i.e. RPI * RHP / HPI.
  * BLS                    CPI rent of primary residence, to extend the US rent index to 2025.
  * US Census Bureau       housing-unit estimates, components of population change,
                           CPS Table AD-1 (young adults at home), completions.
  * ABS, Statistics Canada dwelling stock for Australia and Canada.
  * World Bank DB2020      days to obtain a construction permit (Doing Business 2020,
                           data year 2019; the series was discontinued in 2021).

Two cleaning steps are applied and reported in the page footnotes:
  * Hypostat stock series have breaks (definition changes, one-off typos). A one-year
    outlier is dropped and interpolated; a persistent jump outside -0.5%..+2.5% a year
    is bridged with the country's median growth. More than two problems drops the country.
  * Census revisions show up as population jumps that natural change and migration do
    not explain. Each is spread back linearly over the preceding decade, which is how
    statistical offices rebuild intercensal estimates. (Europe only: the OECD series for
    the other four countries are already intercensal.)
"""

import argparse
import csv
import io
import json
import re
import statistics
import sys
import time
from datetime import date
from pathlib import Path
from typing import Any

import httpx
import openpyxl
import xlrd

HERE = Path(__file__).resolve().parent
RAW = HERE / "raw_responses"
DATA_JSON = HERE / "data.json"
INDEX_HTML = HERE / "index.html"

# id, Eurostat geo (None outside Europe), ISO3 (OECD), Hypostat row name, display name
COUNTRY_ROWS = [
    ("AT", "AT", "AUT", "Austria", "Austria"),
    ("BE", "BE", "BEL", "Belgium", "Belgium"),
    ("CH", "CH", "CHE", "Switzerland", "Switzerland"),
    ("CZ", "CZ", "CZE", "Czechia", "Czechia"),
    ("DE", "DE", "DEU", "Germany", "Germany"),
    ("DK", "DK", "DNK", "Denmark", "Denmark"),
    ("ES", "ES", "ESP", "Spain", "Spain"),
    ("FI", "FI", "FIN", "Finland", "Finland"),
    ("FR", "FR", "FRA", "France", "France"),
    ("GB", "UK", "GBR", "United Kingdom", "UK"),
    ("IE", "IE", "IRL", "Ireland", "Ireland"),
    ("IT", "IT", "ITA", "Italy", "Italy"),
    ("NL", "NL", "NLD", "Netherlands", "Netherlands"),
    ("NO", "NO", "NOR", "Norway", "Norway"),
    ("PL", "PL", "POL", "Poland", "Poland"),
    ("PT", "PT", "PRT", "Portugal", "Portugal"),
    ("SE", "SE", "SWE", "Sweden", "Sweden"),
    ("SK", "SK", "SVK", "Slovakia", "Slovakia"),
    ("US", None, "USA", "USA", "United States"),
    ("CA", None, "CAN", "Canada", "Canada"),
    ("AU", None, "AUS", "Australia", "Australia"),
    ("JP", None, "JPN", "Japan", "Japan"),
]
COUNTRIES = [dict(zip(("id", "eurostat", "iso3", "hypostat", "name"), row)) for row in COUNTRY_ROWS]
BY_ID = {c["id"]: c for c in COUNTRIES}
EUROPE = [c for c in COUNTRIES if c["eurostat"]]
OTHERS = [c for c in COUNTRIES if not c["eurostat"]]
# Doing Business codes: "USA" and "JPN" are New York and Tokyo alone; these are the national aggregates.
DB_CODES = {c["id"]: c["iso3"] for c in COUNTRIES} | {"US": "US", "JP": "JAP"}

VIEWS = {
    "CZ": {"labelled": ["PL", "SK", "AT", "DE"], "parents": ["CZ", "SK", "PL", "AT", "DE", "ES"]},
    "US": {"labelled": ["CA", "AU", "JP", "GB"], "parents": ["US", "CA", "AU", "JP", "ES", "CZ"]},
}

BASE_YEAR = 2015
PANEL_YEARS = list(range(2010, 2026))
PARENTS_YEARS = list(range(2005, 2026))
CZ_YEARS = list(range(2005, 2027))
US_YEARS = list(range(2010, 2026))
COMPLETIONS_YEARS = list(range(2015, 2025))

# Figures that exist only in publications, each with its source.
MANUAL: dict[str, dict[str, Any]] = {
    "cz_dwellings_2011": {
        # Czech census, 26 March 2011: all dwellings, occupied and unoccupied.
        "value": 4_756_572,
        "url": "https://invenio.nusl.cz/record/204140/files/nusl-204140_1.pdf",
    },
    "ca_dwellings_2011": {
        # Canadian census, 10 May 2011: total private dwellings.
        "value": 14_569_633,
        "url": "https://www12.statcan.gc.ca/census-recensement/2011/dp-pd/hlt-fst/pd-pl/index-eng.cfm",
    },
    "jp_dwellings": {
        # Housing and Land Survey, 1 October, all dwellings including vacant ones (final results, table 2-1).
        "values": {2008: 57_586_000, 2013: 60_629_000, 2018: 62_407_000, 2023: 65_047_000},
        "url": "https://www.stat.go.jp/data/jyutaku/2023/pdf/kihon_gaiyou.pdf",
    },
    "parents": {
        # Share of 25-34-year-olds living with a parent, from censuses and surveys.
        "CA": {
            "values": {2011: 19.7, 2016: 21.4, 2021: 23.1},
            "url": "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=9810013701",
        },
        "AU": {"values": {2011: 12.0, 2021: 13.7}, "url": "https://www.abs.gov.au/census/find-census-data"},
        "JP": {
            "values": {2019: 49.6, 2024: 43.0},
            "url": "https://www.ipss.go.jp/ps-dotai/j/DOTAI9/NSHC09_top.asp",
        },
    },
}

EUROSTAT = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data"
OECD_HOUSE = (
    "https://sdmx.oecd.org/public/rest/data/OECD.ECO.MPD,DSD_AN_HOUSE_PRICES@DF_HOUSE_PRICES,1.0/"
    "{areas}.A.RPI+RHP+HPI.?format=csv&startPeriod=2005"
)
OECD_POP = (
    "https://sdmx.oecd.org/public/rest/data/OECD.ELS.SAE,DSD_POPULATION@DF_POP_HIST,1.0/"
    "{areas}.POP.PS._T._T.?format=csv&startPeriod=2005"
)
HYPOSTAT = (
    "https://hypo.org/sites/default/files/2026-09/Final-Statistical-Tables-Hypostat-2026-to-share_UPDATE-14-09.xlsx"
)
CZSO_COMPLETIONS = "https://data.csu.gov.cz/api/dotaz/v1/data/vybery/STA09AT1?format=CSV"
DB2020 = (
    "https://archive.doingbusiness.org/content/dam/doingBusiness/excel/db2020/"
    "Historical-data---COMPLETE-dataset-with-scores.xlsx"
)
BLS_RENT = "https://api.bls.gov/publicAPI/v1/timeseries/data/CUUR0000SEHA"
CENSUS = "https://www2.census.gov/programs-surveys"
US_HU_2010S = f"{CENSUS}/popest/datasets/2010-2020/housing/HU-EST2020_ALL.csv"
US_HU_2020S = f"{CENSUS}/popest/tables/2020-2025/housing/totals/NST-EST2025-HU.xlsx"
US_POP_2010S = f"{CENSUS}/popest/datasets/2010-2020/state/totals/nst-est2020-alldata.csv"
US_POP_2020S = f"{CENSUS}/popest/datasets/2020-2025/state/totals/NST-EST2025-ALLDATA.csv"
US_AD1 = f"{CENSUS}/demo/tables/families/time-series/adults/ad1.xls"
US_COMPLETIONS = "https://www.census.gov/construction/nrc/xls/comps_cust.xlsx"
ABS_DWELLINGS = "https://data.api.abs.gov.au/rest/data/ABS,RES_DWELL_ST,1.0.0/4.AUS.Q?format=csv"
STATCAN_DWELLINGS = (
    "https://www150.statcan.gc.ca/t1/wds/rest/getDataFromVectorByReferencePeriodRange"
    "?vectorIds=%221545147162%22&startRefPeriod=2016-01-01&endReferencePeriod=2026-12-31"
)

SOURCES = {
    "hypostat": {
        "name": "European Mortgage Federation, Hypostat 2026",
        "url": "https://hypo.org/ecbc/publications/hypostat/",
    },
    "czso": {
        "name": "Czech Statistical Office, completed dwellings (DataStat STA09AT1)",
        "url": "https://csu.gov.cz/produkty/dokoncene-byty-v-obcich",
    },
    "census": {
        "name": "Czech Statistical Office, housing stock in the 2011 census",
        "url": MANUAL["cz_dwellings_2011"]["url"],
    },
    "demo": {
        "name": "Eurostat, population change (demo_gind)",
        "url": "https://ec.europa.eu/eurostat/databrowser/view/demo_gind/default/table",
    },
    "parents": {
        "name": "Eurostat, young people living with their parents (ilc_lvps08)",
        "url": "https://ec.europa.eu/eurostat/databrowser/view/ilc_lvps08/default/table",
    },
    "oecd": {
        "name": "OECD, Analytical house price indicators",
        "url": "https://data-explorer.oecd.org/vis?df[ds]=dsDisseminateFinalDMZ&df[id]=DSD_AN_HOUSE_PRICES%40DF_HOUSE_PRICES&df[ag]=OECD.ECO.MPD",
    },
    "oecd_pop": {
        "name": "OECD, historical population data",
        "url": "https://data-explorer.oecd.org/vis?df[ds]=dsDisseminateFinalDMZ&df[id]=DSD_POPULATION%40DF_POP_HIST&df[ag]=OECD.ELS.SAE",
    },
    "bls": {"name": "US Bureau of Labor Statistics, CPI rent of primary residence", "url": "https://www.bls.gov/cpi/"},
    "us_hu": {
        "name": "US Census Bureau, housing unit estimates",
        "url": "https://www.census.gov/programs-surveys/popest/data/tables.html",
    },
    "us_pop": {
        "name": "US Census Bureau, population estimates and components of change",
        "url": "https://www.census.gov/programs-surveys/popest.html",
    },
    "us_ad1": {
        "name": "US Census Bureau, CPS Table AD-1, young adults living at home",
        "url": "https://www.census.gov/data/tables/time-series/demo/families/adults.html",
    },
    "us_nrc": {
        "name": "US Census Bureau, new residential construction (completions)",
        "url": "https://www.census.gov/construction/nrc/index.html",
    },
    "abs": {
        "name": "Australian Bureau of Statistics, number of residential dwellings",
        "url": "https://www.abs.gov.au/statistics/economy/price-indexes-and-inflation/total-value-dwellings",
    },
    "statcan": {
        "name": "Statistics Canada, housing stock (table 36-10-0688) and 2011 census",
        "url": "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=3610068801",
    },
    "jp_hls": {
        "name": "Statistics Bureau of Japan, Housing and Land Survey 2023",
        "url": MANUAL["jp_dwellings"]["url"],
    },
    "ca_parents": {"name": "Statistics Canada, census table 98-10-0137", "url": MANUAL["parents"]["CA"]["url"]},
    "au_parents": {
        "name": "Australian Bureau of Statistics, 2011 and 2021 censuses",
        "url": MANUAL["parents"]["AU"]["url"],
    },
    "jp_parents": {
        "name": "IPSS, National Survey on Household Changes (2019, 2024)",
        "url": MANUAL["parents"]["JP"]["url"],
    },
    "db": {
        "name": "World Bank, Doing Business 2020 (dealing with construction permits)",
        "url": "https://archive.doingbusiness.org/en/data/exploretopics/dealing-with-construction-permits",
    },
}

USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/126.0 Safari/537.36 playground-housing-squeeze/1.0"
)


class Fetcher:
    def __init__(self, refresh: bool, offline: bool) -> None:
        self.refresh = refresh
        self.offline = offline
        self.client = httpx.Client(timeout=300, follow_redirects=True, headers={"User-Agent": USER_AGENT})

    def get(self, name: str, url: str, expect: bytes | None = None) -> bytes:
        """Fetch with a disk cache. `expect` guards against error pages served with HTTP 200."""
        path = RAW / name
        if path.exists() and not self.refresh:
            return path.read_bytes()
        if self.offline:
            sys.exit(f"cache miss in --offline mode: {name}")
        for attempt in range(4):
            print(f"  fetching {name} ...", file=sys.stderr)
            r = self.client.get(url)
            if r.status_code in (429, 502, 503) and attempt < 3:
                time.sleep(10 * (attempt + 1))
                continue
            r.raise_for_status()
            if expect is not None and not r.content.lstrip(b"\xef\xbb\xbf").startswith(expect):
                raise RuntimeError(f"{name}: unexpected content from {url}: {r.content[:80]!r}")
            RAW.mkdir(exist_ok=True)
            path.write_bytes(r.content)
            return r.content
        raise RuntimeError(f"giving up on {url}")


# ---------------------------------------------------------------- parsing helpers


def jsonstat_rows(j: dict[str, Any]) -> list[tuple[dict[str, str], float]]:
    """Flatten a Eurostat JSON-stat 2.0 dataset into (coordinates, value) pairs."""
    dims, sizes = j["id"], j["size"]
    cats = [sorted(j["dimension"][d]["category"]["index"].items(), key=lambda kv: kv[1]) for d in dims]
    out = []
    for flat, value in j["value"].items():
        n, coords = int(flat), {}
        for d, s, c in reversed(list(zip(dims, sizes, cats))):
            coords[d] = c[n % s][0]
            n //= s
        out.append((coords, float(value)))
    return out


def as_year(cell: Any) -> int | None:
    try:
        y = int(float(str(cell).strip()))
    except (TypeError, ValueError):
        return None
    return y if 1950 <= y <= 2100 else None


def as_float(cell: Any) -> float | None:
    try:
        return float(str(cell).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def hypostat_sheet(wb: Any, sheet: str) -> dict[str, dict[int, float]]:
    rows = list(wb[sheet].iter_rows(values_only=True))
    header = next(r for r in rows if sum(as_year(c) is not None for c in r[1:]) > 5)
    years = {j: as_year(c) for j, c in enumerate(header) if j > 0 and as_year(c) is not None}
    out: dict[str, dict[int, float]] = {}
    for r in rows:
        if not isinstance(r[0], str):
            continue
        name = re.sub(r"[*\d]", "", r[0]).strip()
        vals = {years[j]: float(r[j]) for j in years if j < len(r) and isinstance(r[j], (int, float))}
        if vals:
            out[name] = vals
    return out


def census_row(text: str, sumlev: str = "010") -> dict[str, str]:
    """The national row of a Census population-estimates CSV."""
    return next(r for r in csv.DictReader(io.StringIO(text)) if r["SUMLEV"] == sumlev)


def rnd(v: float | None, nd: int = 2) -> float | None:
    return None if v is None else round(v, nd)


def index_to_base(series: dict[int, float], years: list[int]) -> list[float | None] | None:
    if BASE_YEAR not in series:
        return None
    base = series[BASE_YEAR]
    return [rnd(series[y] / base * 100) if y in series else None for y in years]


def interpolate(points: dict[int, float]) -> dict[int, float]:
    """Fill every year between known points with constant (geometric) growth."""
    years = sorted(points)
    out = dict(points)
    for a, b in zip(years, years[1:]):
        r = (points[b] / points[a]) ** (1 / (b - a))
        for y in range(a + 1, b):
            out[y] = points[a] * r ** (y - a)
    return out


def join_names(ids: list[str]) -> str:
    names = [BY_ID[i]["name"] for i in ids]
    return names[0] if len(names) == 1 else ", ".join(names[:-1]) + " and " + names[-1]


# ---------------------------------------------------------------- loaders


def load_hypostat(fetcher: Fetcher) -> tuple[dict[str, dict[int, float]], dict[str, dict[int, float]]]:
    wb = openpyxl.load_workbook(io.BytesIO(fetcher.get("hypostat_2026.xlsx", HYPOSTAT)), read_only=True, data_only=True)
    stock_rows, compl_rows = hypostat_sheet(wb, "16. TDwe Stk"), hypostat_sheet(wb, "14. House Com")
    stock = {c["id"]: stock_rows.get(c["hypostat"], {}) for c in COUNTRIES}
    completions = {c["id"]: compl_rows.get(c["hypostat"], {}) for c in COUNTRIES}
    return stock, completions


def load_population(fetcher: Fetcher) -> dict[str, dict[str, dict[int, float]]]:
    geos = "".join(f"&geo={c['eurostat']}" for c in EUROPE)
    url = f"{EUROSTAT}/demo_gind?indic_de=JAN&indic_de=NATGROW&indic_de=CNMIGRAT&sinceTimePeriod=2000{geos}"
    j = json.loads(fetcher.get("eurostat_demo_gind.json", url))
    geo_to_id = {c["eurostat"]: c["id"] for c in EUROPE}
    out: dict[str, dict[str, dict[int, float]]] = {c["id"]: {"JAN": {}, "NATGROW": {}, "CNMIGRAT": {}} for c in EUROPE}
    for coords, v in jsonstat_rows(j):
        out[geo_to_id[coords["geo"]]][coords["indic_de"]][int(coords["time"])] = v
    return out


def load_oecd_population(fetcher: Fetcher) -> dict[str, dict[int, float]]:
    url = OECD_POP.format(areas="+".join(c["iso3"] for c in OTHERS))
    text = fetcher.get("oecd_population.csv", url, expect=b"DATAFLOW").decode("utf-8-sig")
    iso_to_id = {c["iso3"]: c["id"] for c in OTHERS}
    out: dict[str, dict[int, float]] = {c["id"]: {} for c in OTHERS}
    for r in csv.DictReader(io.StringIO(text)):
        if r["REF_AREA"] in iso_to_id and r["OBS_VALUE"]:
            out[iso_to_id[r["REF_AREA"]]][int(r["TIME_PERIOD"])] = float(r["OBS_VALUE"])
    return out


def load_parents(fetcher: Fetcher) -> dict[str, dict[int, float]]:
    european = sorted({g for v in VIEWS.values() for g in v["parents"] if BY_ID[g]["eurostat"]})
    geos = "".join(f"&geo={g}" for g in european + ["EU27_2020"])
    url = f"{EUROSTAT}/ilc_lvps08?age=Y25-34&sex=T&unit=PC{geos}"
    j = json.loads(fetcher.get("eurostat_ilc_lvps08.json", url))
    out: dict[str, dict[int, float]] = {}
    for coords, v in jsonstat_rows(j):
        out.setdefault(coords["geo"], {})[int(coords["time"])] = v
    return out


def load_us_parents(fetcher: Fetcher) -> dict[int, float]:
    """CPS Table AD-1: share of 25-34-year-olds who are a child of the householder, both sexes."""
    book = xlrd.open_workbook(file_contents=fetcher.get("census_ad1.xls", US_AD1))
    sh = book.sheet_by_index(0)
    start = next(i for i in range(sh.nrows) if str(sh.cell_value(i, 0)).strip().lstrip(".").startswith("25 to 34"))
    out: dict[int, float] = {}
    revised: set[int] = set()
    for i in range(start + 1, sh.nrows):
        m = re.match(r"^\.\.(\d{4})([a-z]*)", str(sh.cell_value(i, 0)).strip())
        if not m:
            break
        year, suffix = int(m.group(1)), m.group(2)
        men_total, men_home = as_float(sh.cell_value(i, 1)), as_float(sh.cell_value(i, 2))
        women_total, women_home = as_float(sh.cell_value(i, 5)), as_float(sh.cell_value(i, 6))
        if None in (men_total, men_home, women_total, women_home):
            continue
        if year in revised:
            continue  # keep the revised ('r') row when a year appears twice
        out[year] = (men_home + women_home) / (men_total + women_total) * 100
        if suffix == "r":
            revised.add(year)
    return out


def load_oecd(fetcher: Fetcher) -> dict[str, dict[str, dict[int, float]]]:
    url = OECD_HOUSE.format(areas="+".join(c["iso3"] for c in COUNTRIES))
    text = fetcher.get("oecd_house_prices.csv", url, expect=b"DATAFLOW").decode("utf-8-sig")
    iso_to_id = {c["iso3"]: c["id"] for c in COUNTRIES}
    out: dict[str, dict[str, dict[int, float]]] = {c["id"]: {"RPI": {}, "RHP": {}, "HPI": {}} for c in COUNTRIES}
    for r in csv.DictReader(io.StringIO(text)):
        if r["REF_AREA"] in iso_to_id and r["OBS_VALUE"]:
            out[iso_to_id[r["REF_AREA"]]][r["MEASURE"]][int(r["TIME_PERIOD"])] = float(r["OBS_VALUE"])
    return out


def load_bls_rent(fetcher: Fetcher) -> dict[int, float]:
    """Annual averages (period M13) of the CPI rent of primary residence. v1 returns the last three years."""
    j = json.loads(fetcher.get("bls_cuur0000seha.json", BLS_RENT, expect=b"{"))
    return {int(d["year"]): float(d["value"]) for d in j["Results"]["series"][0]["data"] if d["period"] == "M13"}


def load_czso_completions(fetcher: Fetcher) -> dict[int, float]:
    text = fetcher.get("czso_sta09at1.csv", CZSO_COMPLETIONS).decode("utf-8-sig")
    return {
        int(r["Roky"]): float(r["Hodnota"])
        for r in csv.DictReader(io.StringIO(text))
        if r["Ukazatel"] == "Dokončené byty" and r["Hodnota"]
    }


def load_us_stock(fetcher: Fetcher) -> dict[int, float]:
    """Census housing-unit estimates on 1 July: the 2010s vintage, then Vintage 2025 from 2020."""
    row = census_row(fetcher.get("census_hu_2010s.csv", US_HU_2010S, expect=b"SUMLEV").decode("latin-1"))
    out = {y: float(row[f"HUESTIMATE{y}"]) for y in range(2010, 2020)}
    wb = openpyxl.load_workbook(io.BytesIO(fetcher.get("census_hu_2020s.xlsx", US_HU_2020S)), read_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    years = next(r for r in rows if r[0] is None and sum(as_year(c) is not None for c in r) >= 5)
    us = next(r for r in rows if isinstance(r[0], str) and r[0].strip(". ") == "United States")
    for j, y in enumerate(years):
        if as_year(y) and isinstance(us[j], (int, float)):
            out[as_year(y)] = float(us[j])
    return out


def load_us_components(fetcher: Fetcher) -> tuple[dict[int, float], dict[int, float]]:
    """Natural change and net international migration, each for the year to 1 July."""
    old = census_row(fetcher.get("census_pop_2010s.csv", US_POP_2010S, expect=b"SUMLEV").decode("latin-1"))
    new = census_row(fetcher.get("census_pop_2020s.csv", US_POP_2020S, expect=b"SUMLEV").decode("latin-1"))
    natural = {y: float(old[f"NATURALINC{y}"]) for y in range(2011, 2021)}
    natural |= {y: float(new[f"NATURALCHG{y}"]) for y in range(2021, 2026)}
    migration = {y: float(new[f"INTERNATIONALMIG{y}"]) for y in range(2021, 2026)}
    return natural, migration


def load_us_completions(fetcher: Fetcher) -> dict[int, float]:
    wb = openpyxl.load_workbook(io.BytesIO(fetcher.get("census_comps.xlsx", US_COMPLETIONS)), read_only=True)
    out = {}
    for r in wb["Annual"].iter_rows(values_only=True):
        if as_year(r[0]) and isinstance(r[1], (int, float)):
            out[as_year(r[0])] = float(r[1]) * 1000
    return out


def load_au_stock(fetcher: Fetcher) -> dict[int, float]:
    """ABS number of residential dwellings, December quarter of each year."""
    text = fetcher.get("abs_res_dwell_st.csv", ABS_DWELLINGS, expect=b"DATAFLOW").decode("utf-8-sig")
    out = {}
    for r in csv.DictReader(io.StringIO(text)):
        if r["TIME_PERIOD"].endswith("-Q4") and r["OBS_VALUE"]:
            out[int(r["TIME_PERIOD"][:4])] = float(r["OBS_VALUE"]) * 10 ** int(r["UNIT_MULT"] or 0)
    return out


def load_ca_stock(fetcher: Fetcher) -> dict[int, float]:
    """StatCan housing stock, second quarter (census-aligned), from 2016; 2011 census before that."""
    j = json.loads(fetcher.get("statcan_36100688.json", STATCAN_DWELLINGS, expect=b"["))
    points = {
        int(p["refPer"][:4]): float(p["value"])
        for p in j[0]["object"]["vectorDataPoint"]
        if p["refPer"][5:7] == "04" and p["value"] is not None
    }
    points[2011] = float(MANUAL["ca_dwellings_2011"]["value"])
    return interpolate(points)


def load_permit_days(fetcher: Fetcher) -> dict[str, float]:
    wb = openpyxl.load_workbook(
        io.BytesIO(fetcher.get("db2020_historical.xlsx", DB2020)), read_only=True, data_only=True
    )
    rows = list(wb["All Data"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] == "Country code")
    header = rows[hi]
    start = header.index("Rank-Dealing with construction permits")
    col = next(j for j in range(start, len(header)) if header[j] == "Time (days)")
    code_to_id = {code: cid for cid, code in DB_CODES.items()}
    out = {}
    for r in rows[hi + 1 :]:
        if r[0] in code_to_id and as_year(r[4]) == 2020:
            v = as_float(r[col])  # the workbook stores numbers as text
            if v is not None:
                out[code_to_id[r[0]]] = v
    return out


# ---------------------------------------------------------------- cleaning


STOCK_LO, STOCK_HI = -0.005, 0.025


def plausible(a: float, b: float, years: int = 1) -> bool:
    return STOCK_LO * years <= b / a - 1 <= STOCK_HI * years


def clean_stock(raw: dict[int, float], years: list[int]) -> tuple[dict[int, float] | None, list[str]]:
    """Repair a Hypostat dwelling-stock series over `years`. Returns (series, fixes) or (None, reason)."""
    v: dict[int, float | None] = {y: raw.get(y) for y in years}
    fixes: list[str] = []
    for a, b, c in zip(years, years[1:], years[2:]):
        va, vb, vc = v[a], v[b], v[c]
        if None not in (va, vb, vc) and not plausible(va, vb) and not plausible(vb, vc) and plausible(va, vc, 2):
            v[b] = None
            fixes.append(f"{b} outlier")
    known = [y for y in years if v[y] is not None]
    if not known:
        return None, ["no data"]
    for a, b in zip(known, known[1:]):
        gap = b - a - 1
        if gap > 2:
            return None, [f"no data {a + 1}–{b - 1}"]
        if gap:
            r = (v[b] / v[a]) ** (1 / (b - a))
            for k in range(1, gap + 1):
                v[a + k] = v[a] * r**k
            fixes.append(f"{a + 1} gap" if gap == 1 else f"{a + 1}–{b - 1} gap")
    span = list(range(known[0], known[-1] + 1))
    growth = {b: v[b] / v[a] - 1 for a, b in zip(span, span[1:])}
    bad = [y for y, g in growth.items() if not STOCK_LO <= g <= STOCK_HI]
    if len(fixes) + len(bad) > 2:
        return None, [f"{len(fixes) + len(bad)} series breaks"]
    if bad:
        med = statistics.median(g for y, g in growth.items() if y not in bad)
        for y in bad:
            growth[y] = med
            fixes.append(f"{y} break")
    chained = {span[0]: v[span[0]]}
    for y in span[1:]:
        chained[y] = chained[y - 1] * (1 + growth[y])
    return chained, fixes


def smooth_census_breaks(pop: dict[str, dict[int, float]]) -> tuple[dict[int, float], list[tuple[int, float]]]:
    """Spread unexplained population jumps (census revisions) back over the preceding decade."""
    jan, nat, mig = pop["JAN"], pop["NATGROW"], pop["CNMIGRAT"]
    out = dict(jan)
    breaks = []
    for y in sorted(jan):
        if all(y - 1 in s for s in (jan, nat, mig)):
            resid = jan[y] - jan[y - 1] - nat[y - 1] - mig[y - 1]
            if abs(resid) > 0.002 * jan[y]:
                breaks.append((y, resid))
                for k in range(1, 10):
                    if y - 10 + k in out:
                        out[y - 10 + k] += resid * k / 10
    return out, breaks


def build_cz_stock(completions: dict[int, float]) -> dict[int, float]:
    """End-of-year Czech dwelling stock: the March 2011 census, plus or minus completions since."""
    stock = {2010: float(MANUAL["cz_dwellings_2011"]["value"])}
    for y in range(2011, max(completions) + 1):
        stock[y] = stock[y - 1] + completions[y]
    for y in range(2010, 2003, -1):
        stock[y - 1] = stock[y] - completions[y]
    return stock


def without_migration(population: dict[int, float], natural: dict[int, float], years: list[int]) -> dict[int, float]:
    """The base-year population carried forward and back by natural change alone.

    `natural[y]` is the change during the step that ends at year y."""
    out = {BASE_YEAR: population[BASE_YEAR]}
    for y in range(BASE_YEAR + 1, years[-1] + 1):
        out[y] = out[y - 1] + natural[y]
    for y in range(BASE_YEAR, years[0], -1):
        out[y - 1] = out[y] - natural[y]
    return out


# ---------------------------------------------------------------- assembly


def panel(pid: str, title: str, series: dict[str, dict[int, float]], notes: list[str]) -> dict[str, Any]:
    out, missing = {}, []
    for c in COUNTRIES:
        s = series.get(c["id"])
        idx = index_to_base(s, PANEL_YEARS) if s else None
        if idx is None:
            missing.append(c["id"])
        else:
            out[c["id"]] = idx
    if missing:
        notes = [f"{join_names(missing)} {'is' if len(missing) == 1 else 'are'} omitted for lack of data."] + notes
    return {"id": pid, "title": title, "series": out, "notes": notes}


def build(fetcher: Fetcher) -> tuple[dict[str, Any], dict[str, Any]]:
    stock_raw, compl_raw = load_hypostat(fetcher)
    pop = load_population(fetcher)
    pop_other = load_oecd_population(fetcher)
    parents = load_parents(fetcher)
    us_parents = load_us_parents(fetcher)
    oecd = load_oecd(fetcher)
    bls_rent = load_bls_rent(fetcher)
    cz_compl = load_czso_completions(fetcher)
    us_compl = load_us_completions(fetcher)
    us_stock = load_us_stock(fetcher)
    us_natural, us_migration = load_us_components(fetcher)
    au_stock = load_au_stock(fetcher)
    ca_stock = load_ca_stock(fetcher)
    jp_stock = interpolate(MANUAL["jp_dwellings"]["values"])
    permits = load_permit_days(fetcher)

    # Chart 1, panel 1: housing stock.
    cz_stock = build_cz_stock(cz_compl)
    stock = {"CZ": cz_stock, "US": us_stock, "CA": ca_stock, "AU": au_stock, "JP": jp_stock}
    stock_fixes = []
    for c in EUROPE:
        if c["id"] in stock:
            continue
        cleaned, fixes = clean_stock(stock_raw[c["id"]], PANEL_YEARS)
        if cleaned is not None:
            stock[c["id"]] = cleaned
            if fixes:
                stock_fixes.append(f"{c['name']} ({', '.join(fixes)})")
    stock_notes = [
        "Czechia is the 2011 census count plus dwellings completed since, without subtracting demolitions.",
        "Canada is interpolated between the 2011 and 2016 censuses, and Japan between its 2013, 2018 and 2023 "
        "housing surveys, which are the latest. Australia starts in 2011.",
    ]
    if stock_fixes:
        stock_notes.append("Series breaks are bridged with typical growth: " + "; ".join(stock_fixes) + ".")

    # Chart 1, panel 2: population, with European census revisions smoothed.
    pop_adj, census_breaks = {}, {}
    for c in EUROPE:
        pop_adj[c["id"]], census_breaks[c["id"]] = smooth_census_breaks(pop[c["id"]])
    pop_adj |= pop_other
    pop_raw = {c["id"]: pop[c["id"]]["JAN"] for c in EUROPE} | pop_other
    in_window = [
        f"{BY_ID[i]['name']} {y}"
        for i, bs in census_breaks.items()
        for y, _ in bs
        if PANEL_YEARS[0] < y <= PANEL_YEARS[-1]
    ]
    pop_notes = ["Population on 1 January in Europe, mid-year in the US, Canada and Australia, and 1 October in Japan."]
    if in_window:
        pop_notes.append(
            "European census revisions (" + ", ".join(in_window) + ") are spread back over the preceding decade."
        )

    # Chart 1, panels 3 and 4: real rents and real house prices.
    us_rpi = oecd["US"]["RPI"]
    if 2025 not in us_rpi and {2024, 2025} <= bls_rent.keys():
        us_rpi[2025] = us_rpi[2024] * bls_rent[2025] / bls_rent[2024]
    real_rents, real_prices = {}, {}
    for c in COUNTRIES:
        o = oecd[c["id"]]
        real_rents[c["id"]] = {
            y: o["RPI"][y] * o["RHP"][y] / o["HPI"][y] for y in o["RHP"] if y in o["RPI"] and y in o["HPI"]
        }
        real_prices[c["id"]] = o["RHP"]
    rent_notes = [
        "Rent indices are deflated by consumer prices. Official rent indices cover all tenancies, including leases "
        "signed long ago, so they rise more slowly than asking rents for new lets.",
        "Czech controlled rents on older leases were phased out between 2007 and 2012. The US 2025 figure extends "
        "the OECD index with BLS data; Australia ends in 2024.",
    ]

    chart1 = {
        "years": PANEL_YEARS,
        "panels": [
            panel("stock", "Housing stock", stock, stock_notes),
            panel("population", "Population", pop_adj, pop_notes),
            panel("rents", "Real rents", real_rents, rent_notes),
            panel("prices", "Real house prices", real_prices, ["House prices deflated by consumer prices."]),
        ],
    }

    # Chart 2: living with parents.
    parents_by_id = {c["id"]: parents.get(c["eurostat"], {}) for c in EUROPE}
    parents_by_id["US"] = us_parents
    for cid, m in MANUAL["parents"].items():
        parents_by_id[cid] = m["values"]
    shown = sorted({g for v in VIEWS.values() for g in v["parents"]})
    chart2 = {
        "years": PARENTS_YEARS,
        "series": {g: [rnd(parents_by_id[g].get(y), 1) for y in PARENTS_YEARS] for g in shown},
        "eu": [rnd(parents.get("EU27_2020", {}).get(y), 1) for y in PARENTS_YEARS],
        "sparse": sorted(MANUAL["parents"]),
    }

    # Chart 3, Czechia: stock vs population vs population without migration, all at 1 January.
    cz = pop["CZ"]
    cz_pop = pop_adj["CZ"]
    cz_natural = without_migration(cz_pop, {y: cz["NATGROW"][y - 1] for y in CZ_YEARS[1:]}, CZ_YEARS)
    cz_stock_jan = {y: cz_stock[y - 1] for y in CZ_YEARS if y - 1 in cz_stock}
    cz_break = dict(census_breaks["CZ"])
    chart3_cz = {
        "years": CZ_YEARS,
        "date": "1 January",
        "population": index_to_base(cz_pop, CZ_YEARS),
        "natural": index_to_base(cz_natural, CZ_YEARS),
        "stock": index_to_base(cz_stock_jan, CZ_YEARS),
        "annotations": [
            {"kind": "marker", "year": 2021, "label": "2021 census"},
            {
                "kind": "arrow",
                "from": 2022,
                "to": 2023,
                "label": "Refugees from Ukraine, 2022",
                "short": "Refugees from Ukraine",
            },
        ],
        "census2021": round(cz_break.get(2021, 0)),
        "refugeeJump": round(cz["JAN"][2023] - cz["JAN"][2022]),
        "naturalChange2025": round(cz["NATGROW"][2025]),
        "sources": ["census", "czso", "demo"],
    }

    # Chart 3, United States: the same three lines, all at 1 July.
    us_pop = pop_other["US"]
    us_nat = without_migration(us_pop, us_natural, US_YEARS)
    chart3_us = {
        "years": US_YEARS,
        "date": "1 July",
        "population": index_to_base(us_pop, US_YEARS),
        "natural": index_to_base(us_nat, US_YEARS),
        "stock": index_to_base(us_stock, US_YEARS),
        "annotations": [
            {
                "kind": "arrow",
                "from": 2022,
                "to": 2024,
                "label": "Immigration surge, 2022–24",
                "short": "Immigration surge",
            },
        ],
        "migration": {str(y): round(v) for y, v in us_migration.items()},
        "naturalChange2025": round(us_natural[2025]),
        "sources": ["us_hu", "us_pop", "oecd_pop"],
    }

    # Chart 4: completions per 1,000 residents and permit days.
    completion_sources = {c["id"]: compl_raw[c["id"]] for c in COUNTRIES} | {"CZ": cz_compl, "US": us_compl}
    # National Canadian completions stop in 2022; Hypostat's later figures repeat earlier years.
    completion_sources["CA"] = {y: v for y, v in compl_raw["CA"].items() if y <= 2022}
    completions = []
    for c in COUNTRIES:
        src, denom = completion_sources[c["id"]], pop_raw[c["id"]]
        rates = [src[y] / denom[y] * 1000 for y in COMPLETIONS_YEARS if y in src and y in denom]
        if len(rates) >= 5:
            completions.append({"id": c["id"], "value": rnd(sum(rates) / len(rates)), "years": len(rates)})
    completions.sort(key=lambda d: -d["value"])
    partial = [d["id"] for d in completions if d["years"] < len(COMPLETIONS_YEARS)]
    compl_missing = [c["id"] for c in COUNTRIES if c["id"] not in {d["id"] for d in completions} and c["id"] != "JP"]
    permit_rows = sorted(({"id": i, "value": v} for i, v in permits.items()), key=lambda d: -d["value"])
    chart4 = {
        "completions": completions,
        "completionsNotes": (
            ([f"{join_names(partial)}: average of the years available."] if partial else [])
            + ([f"{join_names(compl_missing)} omitted for lack of data."] if compl_missing else [])
            + ["Japan publishes housing starts, not completions, so it is not shown."]
        ),
        "permits": permit_rows,
    }

    views = {
        vid: {"focus": vid, "labelled": v["labelled"], "parents": v["parents"]}
        | {"chart3": chart3_cz if vid == "CZ" else chart3_us}
        for vid, v in VIEWS.items()
    }
    payload = {
        "meta": {"built": date.today().isoformat(), "sources": SOURCES},
        "countries": {c["id"]: c["name"] for c in COUNTRIES},
        "views": views,
        "chart1": chart1,
        "chart2": chart2,
        "chart4": chart4,
    }
    checks = {
        "cz_pop_2015_raw": cz["JAN"][2015],
        "cz_census_resid_2021": cz_break.get(2021, 0),
        "cz_parents_2015": parents["CZ"][2015],
        "cz_parents_2025": parents["CZ"][2025],
        "cz_completions_2024": cz_compl[2024],
        "cz_completions_per_1000_2024": cz_compl[2024] / cz["JAN"][2024] * 1000,
        "cz_real_rent_2025": real_rents["CZ"][2025] / real_rents["CZ"][BASE_YEAR] * 100,
        "cz_real_price_2025": real_prices["CZ"][2025] / real_prices["CZ"][BASE_YEAR] * 100,
        "cz_permit_days": permits.get("CZ"),
        "us_pop_2015": us_pop[2015],
        "us_parents_2025": us_parents.get(2025),
        "us_completions_2024": us_compl.get(2024),
        "us_permit_days": permits.get("US"),
        "us_stock_growth_2015_2025": (us_stock[2025] / us_stock[2015] - 1) * 100,
        "us_real_price_2025": real_prices["US"][2025] / real_prices["US"][BASE_YEAR] * 100,
        "us_natural_2025": us_nat[2025] / us_nat[BASE_YEAR] * 100,
    }
    return payload, checks


ANCHORS = {
    "cz_pop_2015_raw": (10_538_275, 10_538_275),
    "cz_census_resid_2021": (-220_000, -195_000),
    "cz_parents_2015": (34.4, 34.4),
    "cz_parents_2025": (22.3, 22.3),
    "cz_completions_2024": (30_274, 30_274),
    "cz_completions_per_1000_2024": (2.6, 3.0),
    "cz_real_rent_2025": (92, 99),
    "cz_real_price_2025": (150, 165),
    "cz_permit_days": (246, 246),
    "us_pop_2015": (321_815_121, 321_815_121),
    "us_parents_2025": (16.4, 16.5),
    "us_completions_2024": (1_626_800, 1_627_000),
    "us_permit_days": (80.5, 80.7),
    "us_stock_growth_2015_2025": (9, 10.5),
    "us_real_price_2025": (150, 158),
    "us_natural_2025": (101.5, 102.8),
}


def validate(checks: dict[str, Any]) -> list[str]:
    failures = []
    for k, (lo, hi) in ANCHORS.items():
        v = checks.get(k)
        if v is None or not lo - 1e-9 <= v <= hi + 1e-9:
            failures.append(f"{k} = {v}, expected {lo}..{hi}")
    return failures


def inline_into_html(payload: dict[str, Any]) -> bool:
    if not INDEX_HTML.exists():
        print(f"  WARNING: {INDEX_HTML} does not exist yet - skipping inline", file=sys.stderr)
        return False
    text = INDEX_HTML.read_text()
    blob = json.dumps(payload, separators=(",", ":"), ensure_ascii=True).replace("<", "\\u003c")
    js = "    const DATA = " + blob + ";\n"
    pattern = re.compile(r"(// DATA_START\n)(.*?)(    // DATA_END)", re.DOTALL)
    if not pattern.search(text):
        print(f"  WARNING: no DATA_START/DATA_END sentinels in {INDEX_HTML}", file=sys.stderr)
        return False
    INDEX_HTML.write_text(pattern.sub(lambda m: m.group(1) + js + m.group(3), text))
    return True


def summarize(payload: dict[str, Any], checks: dict[str, Any]) -> None:
    c1 = payload["chart1"]
    last = c1["years"].index(2025)
    print("\nChart 1 (2015 = 100) in 2025: rank among countries with data, and the median")
    for p in c1["panels"]:
        ranked = sorted(((s[last], k) for k, s in p["series"].items() if s[last] is not None), reverse=True)
        order = [k for _, k in ranked]
        vals = [v for v, _ in ranked]
        parts = []
        for f in ("CZ", "US"):
            v = p["series"].get(f, [None] * len(c1["years"]))[last]
            parts.append(f"{f} {v} (#{order.index(f) + 1})" if f in order else f"{f} {v}")
        print(f"  {p['title']:<18} {'  '.join(parts)}  median {statistics.median(vals):.1f}  n={len(vals)}")
        print(f"      top: {', '.join(f'{k} {v:.0f}' for v, k in ranked[:5])}")
        for n in p["notes"]:
            print(f"      note: {n}")
    c2 = payload["chart2"]
    print("\nChart 2 (% aged 25-34 living with parents), 2015 -> latest:")
    for g, s in c2["series"].items():
        pts = [(y, v) for y, v in zip(c2["years"], s) if v is not None]
        print(f"  {g}: {dict(pts).get(2015)} -> {pts[-1] if pts else None}")
    for vid, view in payload["views"].items():
        c3 = view["chart3"]
        print(f"\nChart 3 {vid} ({c3['date']}, 2015 = 100):")
        for y in c3["years"][:: max(1, len(c3["years"]) // 8)] + [c3["years"][-1]]:
            i = c3["years"].index(y)
            print(f"  {y}: population {c3['population'][i]}  no-migration {c3['natural'][i]}  stock {c3['stock'][i]}")
    print(f"  US net international migration: {payload['views']['US']['chart3']['migration']}")
    c4 = payload["chart4"]
    print("\nChart 4: completions per 1,000 (2015-2024 avg) | permit days (DB2020)")
    print("  " + ", ".join(f"{d['id']} {d['value']}" for d in c4["completions"]))
    print("  " + ", ".join(f"{d['id']} {d['value']:.0f}" for d in c4["permits"]))
    for n in c4["completionsNotes"]:
        print(f"      note: {n}")
    print("\nAnchors:")
    for k, v in checks.items():
        print(f"  {k}: {v:,.2f}" if isinstance(v, float) else f"  {k}: {v}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--refresh", action="store_true", help="ignore the cache and re-download")
    ap.add_argument("--offline", action="store_true", help="fail on a cache miss instead of fetching")
    args = ap.parse_args()
    payload, checks = build(Fetcher(args.refresh, args.offline))
    summarize(payload, checks)
    failures = validate(checks)
    if failures:
        print("\nVALIDATION FAILED:\n  " + "\n  ".join(failures), file=sys.stderr)
        sys.exit(1)
    DATA_JSON.write_text(json.dumps(payload, indent=1, ensure_ascii=False) + "\n")
    inlined = inline_into_html(payload)
    print(f"\nwrote {DATA_JSON.name}" + (" and inlined into index.html" if inlined else ""))


if __name__ == "__main__":
    main()
